"""Core document and evidence workflow. All demo content is synthetic."""
from __future__ import annotations

import io
import json
import os
import re
import threading
import time
import urllib.request
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent
SAMPLES = ROOT / "samples"
_USAGE_LOCK = threading.Lock()
_USAGE_EVENTS: list[dict] = []


def pop_usage_events() -> list[dict]:
    """Return completed API usage records without retaining prompts or responses."""
    with _USAGE_LOCK:
        events = list(_USAGE_EVENTS)
        _USAGE_EVENTS.clear()
    return events

STOP = set("提供提交具备支持包含包括进行以及相关项目服务平台系统方案功能要求投标人供应商必须应须的与和及或在内了" )
ALIASES = {
    "7*24": "7×24", "7x24": "7×24", "24小时全天": "7×24", "营业执照复印件": "营业执照",
    "数据备份方案": "每日备份", "权限": "访问权限", "案例": "项目案例", "元数据": "元数据管理",
}


def normalize(text: str) -> str:
    text = text.lower().replace("７", "7").replace("×", "x").replace("＊", "*")
    for old, new in ALIASES.items():
        text = text.replace(old, new.lower().replace("×", "x"))
    return re.sub(r"\s+", "", text)


def terms(text: str) -> set[str]:
    text = normalize(text)
    latin = re.findall(r"[a-z]+|\d+(?:\.\d+)?", text)
    chinese = re.findall(r"[\u4e00-\u9fff]+", text)
    grams = {part[i:i + 2] for part in chinese for i in range(len(part) - 1)}
    grams |= {part[i:i + 3] for part in chinese for i in range(len(part) - 2)}
    return {x for x in grams | set(latin) if x not in STOP}


def looks_like_tender(text: str) -> bool:
    """Detect requirement documents accidentally uploaded as bidder evidence."""
    normalized = normalize(text)
    markers = ("采购单位", "项目编号", "供应商资格要求", "投标人须", "投标文件须于", "招标文件")
    return sum(marker in normalized for marker in markers) >= 2


def read_document(filename: str, blob: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in (".txt", ".md"):
        return blob.decode("utf-8-sig")
    if suffix == ".pdf":
        from pypdf import PdfReader
        pages = PdfReader(io.BytesIO(blob)).pages
        return "\n".join(page.extract_text() or "" for page in pages)
    if suffix == ".docx":
        from docx import Document
        document = Document(io.BytesIO(blob))
        pieces = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            pieces.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(pieces)
    raise ValueError("仅支持 TXT、MD、PDF、DOCX 文件")


def section_for(line: str, current: str) -> str:
    if any(x in line for x in ("资格要求", "供应商资格")):
        return "资格要求"
    if any(x in line for x in ("技术服务要求", "技术要求", "建设要求")):
        return "技术要求"
    if "商务与交付" in line or "服务要求" in line:
        return "商务与交付"
    return current


def extract_requirements(text: str) -> list[dict]:
    rows = []
    section = "其他要求"
    for line_no, raw in enumerate(text.splitlines(), 1):
        line = raw.strip()
        if not line:
            continue
        section = section_for(line, section)
        match = re.match(r"^(?:\d{1,2}[\.、]|\(\d{1,2}\))\s*(.+)", line)
        if not match:
            continue
        body = match.group(1).strip()
        if len(body) < 8 or not re.search(r"须|应|提供|支持|提交|故障|备份|服务", body):
            continue
        category = "时间节点" if re.search(r"\d{4}年\d{1,2}月\d{1,2}日|截止|前提交", body) else section
        rows.append({
            "id": f"R{len(rows) + 1:02d}", "text": body, "source": line,
            "line": line_no, "category": category, "priority": "高" if re.search(r"须|必须|不得|截止", body) else "中",
            "review_status": "pending", "review_note": "",
        })
    return rows


def rank_evidence(query: str, library: list[dict]) -> list[dict]:
    query_terms = terms(query)
    ranked = []
    for item in library:
        if looks_like_tender(item.get("text", "")):
            continue
        evidence_terms = terms(item["name"] + " " + item["text"])
        overlap = query_terms & evidence_terms
        score = sum(2 if len(term) >= 3 else 1 for term in overlap)
        if score >= 3:
            ranked.append({"id": item["id"], "name": item["name"], "type": item["type"],
                           "excerpt": item["text"][:280], "score": score})
    ranked.sort(key=lambda x: (-x["score"], x["name"]))
    return ranked[:3]


def classify(req: dict, evidence: list[dict]) -> tuple[str, str]:
    q = normalize(req["text"])
    if req["category"] == "时间节点":
        return "action", "提交时间需由项目负责人跟进"
    if not evidence:
        return "gap", "材料库未检索到相关证据"
    top = normalize(evidence[0]["excerpt"])
    evidence_texts = [normalize(item["excerpt"]) for item in evidence]
    combined = normalize(" ".join(item["excerpt"] for item in evidence))
    def direct_proof(required: tuple[str, ...], rejected: tuple[str, ...] = ()) -> bool:
        return any(all(word in text for word in required) and not any(word in text for word in rejected)
                   for text in evidence_texts)
    if "不少于三个" in q or "至少三个" in q or "三个政务" in q:
        positive = any("政务" in text and any(mark in text for mark in
                       ("共三个", "累计三个", "已完成三个", "不少于三个", "至少三个", "3个政务", "三个政务"))
                       and not any(mark in text for mark in ("不足三个", "仅有两个", "只有两个"))
                       for text in evidence_texts)
        if positive:
            return "supported", "已找到不少于三个政务项目案例的直接材料"
        return "gap", "现有资料仅列出两个政务项目，数量不足"
    if "至少两个高校" in q or "两个高校" in q:
        positive = any("高校" in text and any(mark in text for mark in
                       ("共两个", "累计两个", "已完成两个", "不少于两个", "至少两个", "2个高校", "两个高校"))
                       and not any(mark in text for mark in ("不足两个", "仅有一个", "只有一个"))
                       for text in evidence_texts)
        if positive:
            return "supported", "已找到不少于两个高校项目案例的直接材料"
        return "gap", "现有资料仅列出一个高校项目，数量不足"
    if "信息系统项目管理师" in q:
        if direct_proof(("信息系统项目管理师", "证书"), ("未取得", "不具备", "待补充", "模板")):
            return "supported", "已找到项目负责人信息系统项目管理师证书材料"
        return "gap", "未找到项目负责人证书"
    if "4小时内恢复" in q:
        if direct_proof(("4小时内恢复",), ("未承诺", "不能承诺", "待确认")):
            return "supported", "已找到重大故障4小时内恢复的服务承诺"
        return "gap", "资料未承诺4小时内恢复，须负责人确认"
    if "不少于一年的运维" in q:
        if any(("不少于一年" in text or "至少一年" in text or "12个月" in text) and "未承诺" not in text
               for text in evidence_texts):
            return "supported", "已找到不少于一年运维服务期限的材料"
        return "gap", "资料未承诺一年运维期限"
    if "实施计划" in q or "迁移与验收方案" in q:
        if direct_proof(("交接", "试运行", "验收"), ("需按", "待填写", "待确认", "模板")):
            return "supported", "已找到包含交接、试运行和验收节点的实施计划"
        return "review", "找到方案模板，项目日期与验收标准仍需补充"
    if "人员清单" in q:
        if any(("姓名" in text or "人员清单" in text) and "岗位职责" in text
               and not any(word in text for word in ("模板", "待填写", "需人工确认", "具体姓名和资质需确认"))
               for text in evidence_texts):
            return "supported", "已找到包含具体人员与岗位职责的团队清单"
        return "review", "找到岗位模板，具体人员与资质仍需确认"
    if "营业执照" in q and "营业执照" in combined:
        return "supported", "已找到营业执照摘要；递交前须核对原件与有效期"
    if "7x24" in q and "7x24" in combined:
        return "supported", "服务渠道与7×24小时承诺有材料依据"
    if "2小时内响应" in q and "2小时内响应" in combined:
        return "supported", "已找到2小时响应承诺"
    if "季度安全巡检" in q and "每季度开展安全巡检" in combined:
        return "supported", "已找到每季度安全巡检和漏洞报告说明"
    signals = ["每日备份", "季度安全巡检", "质量规则", "质量报告", "访问权限控制", "元数据管理", "多源数据接入", "数据库", "api"]
    matched = any(signal in q and signal in combined for signal in signals)
    if matched:
        return "supported", "已找到直接对应的产品或服务说明"
    if len(terms(req["text"]) & terms(evidence[0]["excerpt"])) >= 5:
        return "review", "检索到相近材料，仍需人工核对是否完全满足"
    return "gap", "证据关联较弱，建议补充材料"


def analyze(text: str, library: list[dict]) -> list[dict]:
    result = []
    for req in extract_requirements(text):
        evidence = rank_evidence(req["text"], library)
        status, reason = classify(req, evidence)
        req.update({"evidence": evidence, "status": status, "reason": reason})
        result.append(req)
    return result


def refresh_evidence(requirements: list[dict], library: list[dict]) -> list[dict]:
    """Re-evaluate existing extracted requirements after the material library changes."""
    refreshed = []
    for original in requirements:
        req = dict(original)
        evidence = rank_evidence(req["text"], library)
        req["evidence"] = evidence
        req["status"], req["reason"] = classify(req, evidence)
        refreshed.append(req)
    return refreshed


def _chat_completion(messages: list[dict], max_tokens: int = 1200, json_mode: bool = False) -> str:
    key = os.getenv("LLM_API_KEY")
    if not key:
        raise ValueError("未配置 LLM_API_KEY")
    base = os.getenv("LLM_BASE_URL", "https://api.openai.com/v1").rstrip("/")
    model = os.getenv("LLM_MODEL", "gpt-4o-mini")
    body = {"model": model, "temperature": 0, "max_tokens": max_tokens, "messages": messages}
    if "aliyuncs.com" in base and model.startswith("qwen3.5-"):
        body["enable_thinking"] = False
        if json_mode:
            body["response_format"] = {"type": "json_object"}
    payload = json.dumps(body).encode("utf-8")
    request = urllib.request.Request(base + "/chat/completions", payload,
                                     {"Content-Type": "application/json", "Authorization": "Bearer " + key})
    try:
        timeout = max(5, min(120, int(os.getenv("LLM_TIMEOUT_SECONDS", "90"))))
    except ValueError:
        timeout = 90
    started = time.perf_counter()
    with urllib.request.urlopen(request, timeout=timeout) as response:
        output = json.load(response)
    usage = output.get("usage") or {}
    prompt_tokens = usage.get("prompt_tokens")
    completion_tokens = usage.get("completion_tokens")
    total_tokens = usage.get("total_tokens")
    if total_tokens is None and isinstance(prompt_tokens, int) and isinstance(completion_tokens, int):
        total_tokens = prompt_tokens + completion_tokens
    event = {"at": datetime.now().isoformat(timespec="seconds"), "model": model,
             "prompt_tokens": prompt_tokens if isinstance(prompt_tokens, int) else None,
             "completion_tokens": completion_tokens if isinstance(completion_tokens, int) else None,
             "total_tokens": total_tokens if isinstance(total_tokens, int) else None,
             "latency_ms": round((time.perf_counter() - started) * 1000)}
    with _USAGE_LOCK:
        _USAGE_EVENTS.append(event)
    return output["choices"][0]["message"]["content"].strip()


def _json_array(text: str) -> list:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE)
    value = json.loads(text)
    if isinstance(value, dict) and isinstance(value.get("requirements"), list):
        value = value["requirements"]
    if not isinstance(value, list):
        raise ValueError("模型未返回 JSON 数组")
    return value


def extract_requirements_with_model(text: str) -> list[dict]:
    """Model extracts exact source lines; invalid or invented quotes are discarded."""
    lines = text.splitlines()
    chunks = []
    current = []
    current_size = 0
    for number, raw in enumerate(lines, 1):
        item = f"L{number}: {raw[:1000]}"
        if current and current_size + len(item) > 10000:
            chunks.append(current)
            current, current_size = [], 0
        current.append(item)
        current_size += len(item)
    if current:
        chunks.append(current)
    if len(chunks) > 20:
        raise ValueError("文档过长，模型抽取最多支持约 20 个文本分块")
    extracted = []
    seen_lines = set()
    instruction = (ROOT / "prompts" / "requirement_extraction_v1.txt").read_text(encoding="utf-8")
    for chunk in chunks:
        reply = _chat_completion([{"role": "system", "content": instruction},
                                  {"role": "user", "content": "\n".join(chunk)}], max_tokens=2500, json_mode=True)
        for item in _json_array(reply):
            if not isinstance(item, dict):
                continue
            line_no = item.get("line")
            quote = str(item.get("quote", "")).strip()
            if not isinstance(line_no, int) or not (1 <= line_no <= len(lines)) or line_no in seen_lines:
                continue
            source = lines[line_no - 1].strip()
            if len(quote) < 7 or quote not in source:
                continue
            seen_lines.add(line_no)
            category = item.get("category", "其他要求")
            if category not in ("资格要求", "技术要求", "商务与交付", "时间节点", "其他要求"):
                category = "其他要求"
            extracted.append({"id": "", "text": quote, "source": source, "line": line_no,
                              "category": category, "priority": "高" if re.search(r"须|必须|不得|截止", quote) else "中",
                              "review_status": "pending", "review_note": ""})
    extracted.sort(key=lambda row: row["line"])
    for index, row in enumerate(extracted, 1):
        row["id"] = f"R{index:02d}"
    if not extracted:
        raise ValueError("模型未返回可验证的招标要求")
    return extracted


def analyze_with_model(text: str, library: list[dict]) -> tuple[list[dict], str, str]:
    """Configured model extracts requirements; local tools retrieve and classify evidence."""
    if not os.getenv("LLM_API_KEY"):
        return analyze(text, library), "离线规则分析", ""
    try:
        extracted = extract_requirements_with_model(text)
        for req in extracted:
            evidence = rank_evidence(req["text"], library)
            status, reason = classify(req, evidence)
            req.update({"evidence": evidence, "status": status, "reason": reason})
        return extracted, "模型抽取 + 本地证据检索", ""
    except Exception as exc:
        return analyze(text, library), "离线规则分析", f"模型抽取失败，已自动回退：{type(exc).__name__}: {exc}"


CHAPTER_FIELDS = {
    "project_overview": ("项目概述", ("项目背景", "项目目标", "对招标需求的理解")),
    "technical_service": ("技术与服务方案", ("总体解决方案", "运维/技术服务内容", "故障响应与应急处理", "数据备份与恢复", "安全保障方案")),
    "implementation": ("项目实施方案", ("项目实施阶段", "实施计划", "主要交付物", "验收方案")),
    "team": ("项目团队", ("项目组织架构", "人员清单", "岗位职责", "人员相关资质")),
    "service_quality": ("服务与质量保障", ("服务时间与服务渠道", "SLA响应机制", "质量保障措施", "售后及运维保障")),
}


def _fallback_chapters(project: dict, requirements: list[dict]) -> dict[str, dict[str, str]]:
    supported = [row for row in requirements if row.get("status") == "supported" and row.get("evidence")]

    def grounded(keywords: tuple[str, ...], missing: str) -> str:
        rows = [row for row in supported if any(word in (row.get("text", "") + " " + row["evidence"][0]["excerpt"])
                                                    for word in keywords)]
        if not rows:
            return "【待补充】" + missing
        pieces = []
        used = set()
        for row in rows[:3]:
            evidence = row["evidence"][0]
            if evidence["name"] in used:
                continue
            used.add(evidence["name"])
            excerpt = evidence["excerpt"].replace("【虚构企业材料】", "").strip().rstrip("。；; ")
            pieces.append(f"依据《{evidence['name']}》，现有材料表明：{excerpt}")
        return "；".join(pieces) + "。正式提交前应核验原件及适用范围。"

    total = len(requirements)
    technical = sum(row.get("category") == "技术要求" for row in requirements)
    return {
        "project_overview": {
            "项目背景": f"本项目为“{project.get('name', '投标项目')}”，采购单位为{project.get('customer', '【待补充】')}。本文件依据招标要求及当前企业材料编制。",
            "项目目标": f"围绕招标文件识别的 {total} 项要求形成可追溯响应，其中包含 {technical} 项技术要求，并对材料不足事项保留【待补充】标记。",
            "对招标需求的理解": "本项目需在资格合规、技术服务、实施交付、团队组织和持续保障之间形成完整闭环，并确保每项承诺均有企业材料支持。",
        },
        "technical_service": {
            "总体解决方案": grounded(("平台", "运维", "技术"), "需结合项目范围补充总体架构、服务边界和实施路径。"),
            "运维/技术服务内容": grounded(("运维", "巡检", "服务"), "需补充具体服务目录、作业流程和交付标准。"),
            "故障响应与应急处理": grounded(("故障", "响应", "恢复"), "需补充故障分级、响应时限、升级路径和应急联系人。"),
            "数据备份与恢复": grounded(("备份", "恢复"), "需补充备份频率、保存位置、恢复流程和演练安排。"),
            "安全保障方案": grounded(("安全", "漏洞", "权限"), "需补充安全巡检、权限管理、漏洞整改和事件处置措施。"),
        },
        "implementation": {
            "项目实施阶段": grounded(("实施计划", "迁移", "试运行", "验收"), "需补充启动、交接、实施、试运行和验收阶段。"),
            "实施计划": "【待补充】需结合中标日期、招标节点和资源安排补充任务、负责人、起止时间及里程碑。",
            "主要交付物": "【待补充】需按招标要求确认方案文档、实施记录、巡检报告、培训材料和验收材料清单。",
            "验收方案": grounded(("验收",), "需补充验收范围、验收标准、测试方法、交付资料和双方确认流程。"),
        },
        "team": {
            "项目组织架构": grounded(("团队", "岗位", "人员"), "需补充项目经理、技术、运维、安全及质量角色的组织关系。"),
            "人员清单": "【待补充】需填写姓名、岗位、投入阶段、联系方式及在岗安排。",
            "岗位职责": grounded(("岗位职责",), "需补充各岗位职责、协作边界和升级机制。"),
            "人员相关资质": grounded(("证书", "资质"), "需补充项目负责人及核心成员的证书名称、编号和有效期。"),
        },
        "service_quality": {
            "服务时间与服务渠道": grounded(("服务渠道", "7×24", "7x24"), "需补充服务时间、热线、工单系统和联系人。"),
            "SLA响应机制": grounded(("响应", "恢复", "sla"), "需补充事件分级、响应时限、恢复时限和升级机制。"),
            "质量保障措施": grounded(("质量", "巡检", "报告"), "需补充质量目标、检查频率、报告机制和改进闭环。"),
            "售后及运维保障": grounded(("运维", "售后", "服务"), "需补充服务周期、保障范围、续保方式和服务退出安排。"),
        },
    }


def _llm_bid_content(project: dict, requirements: list[dict]) -> dict:
    if not os.getenv("LLM_API_KEY") or not requirements:
        return {}
    prompt_path = ROOT / "prompts" / "bid_document_generation_v2.txt"
    instruction = prompt_path.read_text(encoding="utf-8")
    payload = {"project": {"name": project.get("name", "投标项目"), "customer": project.get("customer", "待填写")},
               "requirements": []}
    for row in requirements:
        payload["requirements"].append({"id": row["id"], "category": row["category"],
                                        "requirement": row["text"], "source": row["source"],
                                        "status": {"supported": "有依据", "review": "待确认", "gap": "材料缺失", "action": "待确认"}.get(row["status"], "待确认"),
                                        "evidence": [{"name": item["name"], "type": item["type"], "text": item["excerpt"]}
                                                     for item in (row.get("evidence", [])[:2] if row.get("status") == "supported" else [])]})
    try:
        reply = _chat_completion([{"role": "system", "content": instruction},
                                  {"role": "user", "content": json.dumps(payload, ensure_ascii=False)}],
                                 max_tokens=5000, json_mode=True)
        output = json.loads(re.sub(r"^```(?:json)?\s*|\s*```$", "", reply.strip(), flags=re.IGNORECASE))
        if not isinstance(output, dict):
            return {}
        known_numbers = set(re.findall(r"\d+(?:\.\d+)?", json.dumps(payload, ensure_ascii=False)))

        def safe_text(value, limit=1800) -> str:
            value = str(value or "").strip()[:limit]
            claimed = set(re.findall(r"\d+(?:\.\d+)?", value))
            return value if value and claimed <= known_numbers else "【待补充】该内容包含输入材料无法验证的信息，请人工核对。"

        chapters = {}
        raw_chapters = output.get("chapters", {}) if isinstance(output.get("chapters"), dict) else {}
        for key, (_, fields) in CHAPTER_FIELDS.items():
            raw = raw_chapters.get(key, {}) if isinstance(raw_chapters.get(key), dict) else {}
            chapters[key] = {field: safe_text(raw.get(field)) for field in fields if raw.get(field)}
        allowed = {row["id"]: row for row in requirements}
        responses = {}
        for item in output.get("responses", []):
            if not isinstance(item, dict) or item.get("id") not in allowed:
                continue
            req = allowed[item["id"]]
            if req["status"] != "supported" or not req.get("evidence"):
                continue
            text = str(item.get("text") or "").strip()[:800]
            known = set(re.findall(r"\d+(?:\.\d+)?", req["text"] + " ".join(row["excerpt"] for row in req["evidence"][:2])))
            claimed = set(re.findall(r"\d+(?:\.\d+)?", text))
            if text and "【待补充】" not in text and claimed <= known:
                responses[req["id"]] = text
        commercial = output.get("commercial", {}) if isinstance(output.get("commercial"), dict) else {}
        return {"chapters": chapters, "responses": responses,
                "commercial": {"商务条款响应": safe_text(commercial.get("商务条款响应"), 1200),
                               "项目报价": "【待补充】需由商务负责人依据最终范围和报价审批结果填写。"}}
    except Exception:
        return {}


def generate_draft(project: dict, use_model: bool = True) -> dict:
    sections = []
    requirements = project.get("requirements", [])
    model_content = _llm_bid_content(project, requirements) if use_model else {}
    model_responses = model_content.get("responses", {})
    for req in requirements:
        evidence = req.get("evidence", [])
        if req["status"] == "supported" and evidence:
            llm_text = model_responses.get(req["id"])
            excerpt = evidence[0]["excerpt"].replace("【虚构企业材料】", "").strip().rstrip("。；; ")
            response = llm_text if llm_text else f"根据《{evidence[0]['name']}》，投标人现有材料载明：{excerpt}。该项按招标要求响应，正式提交前复核原件及适用范围。"
        elif req["status"] == "review":
            response = f"【待补充】待人工确认：{req['reason']}。参考材料：《{evidence[0]['name']}》。" if evidence else f"【待补充】待人工确认：{req['reason']}。"
        elif req["status"] == "action":
            response = f"【待补充】待项目负责人跟进：{req['text']}"
        else:
            response = f"【待补充】材料缺口：{req['reason']}；材料补齐前暂不作符合性承诺。"
        sections.append({"requirement_id": req["id"], "category": req["category"], "requirement": req["text"],
                         "response": response, "source": req["source"],
                         "evidence_name": evidence[0]["name"] if evidence and req["status"] in ("supported", "review") else "",
                         "status": req["status"],
                         "response_status": {"supported": "完全响应", "review": "待确认", "gap": "材料缺失", "action": "待确认"}.get(req["status"], "待确认")})
    summary = {key: sum(x["status"] == key for x in sections) for key in ("supported", "review", "gap", "action")}
    categories = []
    for category in ("资格要求", "技术要求", "商务与交付", "时间节点", "其他要求"):
        rows = [row for row in sections if row["category"] == category]
        if rows:
            categories.append({"name": category, "count": len(rows),
                               "supported": sum(row["status"] == "supported" for row in rows),
                               "attention": sum(row["status"] != "supported" for row in rows)})
    evidence_names = []
    attachments = []
    case_studies = []
    for req in requirements:
        for evidence in req.get("evidence", [])[:1]:
            if evidence["name"] not in evidence_names:
                evidence_names.append(evidence["name"])
                attachments.append({"name": evidence["name"], "type": evidence["type"],
                                    "description": evidence["excerpt"].replace("【虚构企业材料】", "").strip()})
                if evidence["type"] == "案例":
                    case_studies.append({"name": evidence["name"], "description": evidence["excerpt"].replace("【虚构企业材料】", "").strip(),
                                         "evidence_name": evidence["name"]})
    fallback = _fallback_chapters(project, requirements)
    for key, values in model_content.get("chapters", {}).items():
        if key in fallback:
            fallback[key].update({name: text for name, text in values.items() if text})
    chapters = [{"key": key, "title": title, "items": [{"title": field, "content": fallback[key][field]} for field in fields]}
                for key, (title, fields) in CHAPTER_FIELDS.items()]
    commercial = {"商务条款响应": "【待补充】需由商务负责人逐项确认付款、履约、税费及合同条款。",
                  "项目报价": "【待补充】需由商务负责人依据最终范围和报价审批结果填写。"}
    commercial.update({key: value for key, value in model_content.get("commercial", {}).items() if value})
    material_gaps = [{"requirement_id": row["requirement_id"], "requirement": row["requirement"],
                      "status": row["response_status"], "needed": next(req["reason"] for req in requirements if req["id"] == row["requirement_id"])}
                     for row in sections if row["status"] != "supported"]
    return {"generated_at": datetime.now().isoformat(timespec="seconds"),
            "mode": "模型辅助生成" if model_content else "离线可追溯模板",
            "document_title": project.get("name", "投标项目") + " 投标文件初稿",
            "project_name": project.get("name", "投标项目"),
            "customer": project.get("customer", "待填写"),
            "chapters": chapters, "sections": sections, "categories": categories,
            "case_studies": case_studies, "commercial": commercial, "attachments": attachments,
            "material_gaps": material_gaps, "evidence_catalog": evidence_names,
            "summary": summary,
            "review_summary": {"confirmed": sum(req.get("review_status") == "confirmed" for req in requirements),
                               "needs_work": sum(req.get("review_status") == "needs_work" for req in requirements),
                               "pending": sum(req.get("review_status", "pending") == "pending" for req in requirements)}}


def export_docx(project: dict) -> bytes:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    def shade(cell, fill: str) -> None:
        properties = cell._tc.get_or_add_tcPr()
        element = OxmlElement("w:shd")
        element.set(qn("w:fill"), fill)
        properties.append(element)

    def set_cell_text(cell, value: str, bold: bool = False) -> None:
        cell.text = ""
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(str(value))
        run.bold = bold
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(9)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    document = Document()
    draft = project.get("draft", {})
    sections = draft.get("sections", [])
    labels = {"supported": "有依据", "review": "待确认", "gap": "材料缺口", "action": "待跟进"}
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(10.5)
    for style_name, size, color in (("Title", 28, "173B46"), ("Heading 1", 18, "087C75"), ("Heading 2", 14, "183145")):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.4)
    section.left_margin = section.right_margin = Cm(2.3)

    cover = document.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(100)
    title = cover.add_run(draft.get("document_title") or project["name"] + " 投标文件初稿")
    title.bold = True
    title.font.name = "Microsoft YaHei"
    title.font.size = Pt(28)
    title.font.color.rgb = RGBColor(8, 124, 117)
    subtitle = document.add_paragraph("技术与商务响应文件")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(16)
    subtitle.runs[0].font.color.rgb = RGBColor(70, 97, 107)
    document.add_paragraph("")
    meta = document.add_table(rows=3, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, pair in enumerate((("项目名称", project.get("name", "待填写")),
                                  ("采购单位", project.get("customer", "待填写")),
                                  ("生成时间", draft.get("generated_at", "待填写")))):
        set_cell_text(meta.cell(index, 0), pair[0], True); shade(meta.cell(index, 0), "EAF5F1")
        set_cell_text(meta.cell(index, 1), pair[1])
    warning = document.add_paragraph("内部初稿｜提交前须完成资质原件、服务承诺、报价及签章复核")
    warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
    warning.paragraph_format.space_before = Pt(70)
    warning.runs[0].font.color.rgb = RGBColor(189, 100, 93)
    document.add_page_break()
    chapters = draft.get("chapters", [])
    if not chapters:
        fallback = _fallback_chapters(project, project.get("requirements", []))
        chapters = [{"key": key, "title": title, "items": [{"title": field, "content": fallback[key][field]} for field in fields]}
                    for key, (title, fields) in CHAPTER_FIELDS.items()]
    chapter_map = {chapter["key"]: chapter for chapter in chapters}

    def add_chapter(number: str, key: str) -> None:
        chapter = chapter_map[key]
        document.add_heading(f"{number}、{chapter['title']}", level=1)
        for item in chapter.get("items", []):
            document.add_heading(item["title"], level=2)
            document.add_paragraph(item["content"])

    add_chapter("一", "project_overview")
    add_chapter("二", "technical_service")
    add_chapter("三", "implementation")
    implementation = chapter_map["implementation"].get("items", [])
    plan_table = document.add_table(rows=1, cols=2); plan_table.style = "Table Grid"
    for cell, value in zip(plan_table.rows[0].cells, ("实施要素", "计划内容")):
        set_cell_text(cell, value, True); shade(cell, "DDEFEA")
    for item in implementation:
        row = plan_table.add_row().cells; set_cell_text(row[0], item["title"], True); set_cell_text(row[1], item["content"])

    add_chapter("四", "team")
    team_table = document.add_table(rows=1, cols=4); team_table.style = "Table Grid"
    for cell, value in zip(team_table.rows[0].cells, ("岗位", "人员", "主要职责", "相关资质")):
        set_cell_text(cell, value, True); shade(cell, "DDEFEA")
    row = team_table.add_row().cells
    for cell, value in zip(row, ("项目负责人及核心成员", "【待补充】", "详见本章岗位职责", "【待补充】")):
        set_cell_text(cell, value)

    add_chapter("五", "service_quality")

    document.add_heading("六、类似项目案例", level=1)
    cases = draft.get("case_studies", [])
    if cases:
        for index, case in enumerate(cases, 1):
            document.add_heading(f"案例 {index}：{case['name']}", level=2)
            document.add_paragraph(case["description"])
            document.add_paragraph("证明材料：" + case["evidence_name"])
    else:
        document.add_paragraph("【待补充】当前企业材料中未找到可验证的类似项目案例。")

    document.add_page_break()
    document.add_heading("七、招标要求响应表", level=1)
    matrix = document.add_table(rows=1, cols=5); matrix.style = "Table Grid"; matrix.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(matrix.rows[0].cells, ("招标条款", "招标要求", "投标响应", "证明材料", "响应状态")):
        set_cell_text(cell, value, True); shade(cell, "DDEFEA")
    for item in sections:
        row = matrix.add_row().cells
        for cell, value in zip(row, (item["requirement_id"], item["requirement"], item["response"],
                                     item["evidence_name"] or "【待补充】", item.get("response_status", "待确认"))):
            set_cell_text(cell, value)

    document.add_heading("八、商务响应", level=1)
    commercial = draft.get("commercial", {})
    for title in ("商务条款响应", "项目报价"):
        document.add_heading(title, level=2)
        document.add_paragraph(commercial.get(title) or "【待补充】")

    document.add_heading("九、附件与证明材料", level=1)
    attachments = draft.get("attachments", [])
    if attachments:
        attachment_table = document.add_table(rows=1, cols=3); attachment_table.style = "Table Grid"
        for cell, value in zip(attachment_table.rows[0].cells, ("材料名称", "材料类型", "材料说明")):
            set_cell_text(cell, value, True); shade(cell, "DDEFEA")
        for item in attachments:
            row = attachment_table.add_row().cells
            for cell, value in zip(row, (item["name"], item["type"], item["description"])):
                set_cell_text(cell, value)
    else:
        document.add_paragraph("【待补充】营业执照、资质证书、人员证书、案例证明及其他招标要求附件。")

    document.add_heading("十、材料缺口清单", level=1)
    gaps = draft.get("material_gaps", [])
    if gaps:
        gap_table = document.add_table(rows=1, cols=4); gap_table.style = "Table Grid"
        for cell, value in zip(gap_table.rows[0].cells, ("条款", "响应状态", "待补充内容", "需要补充的材料")):
            set_cell_text(cell, value, True); shade(cell, "FBE9E6")
        for item in gaps:
            row = gap_table.add_row().cells
            for cell, value in zip(row, (item["requirement_id"], item["status"], item["requirement"], item["needed"])):
                set_cell_text(cell, value)
    else:
        document.add_paragraph("当前未识别到材料缺口；正式提交前仍须人工核验全部原件、报价及承诺。")
    document.add_paragraph("本文件由系统辅助生成，属于内部编辑初稿，不代表已通过资格审查或形成正式承诺。")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def evaluate(library: list[dict]) -> dict:
    """Evaluate the deterministic baseline against independently labelled files."""
    started = time.perf_counter()
    tasks = []
    for filename in ("gold.json", "gold_extended.json"):
        tasks.extend(json.loads((SAMPLES / filename).read_text(encoding="utf-8"))["projects"])
    counts = {key: 0 for key in ("gold", "found", "status_correct", "evidence_correct",
                                  "evidence_evaluated", "end_to_end_correct", "draft_guardrail_correct",
                                  "extracted", "predicted_supported", "false_supported")}
    errors = []
    per_project = []
    per_split = {}
    for task in tasks:
        split = task.get("split", "未分类")
        group = per_split.setdefault(split, {key: 0 for key in counts})
        project_counts = {key: 0 for key in counts}
        text = (SAMPLES / task["file"]).read_text(encoding="utf-8")
        predictions = analyze(text, library)
        draft = generate_draft({"requirements": predictions}, use_model=False)
        sections = {row["requirement_id"]: row for row in draft["sections"]}
        used = set()
        for label in task["gold"]:
            project_counts["gold"] += 1
            matched = next((r for r in predictions if label["needle"] in r["text"] and r["id"] not in used), None)
            if matched is None:
                errors.append({"file": task["file"], "needle": label["needle"], "type": "要求遗漏", "detail": "未提取到该要求"})
                continue
            used.add(matched["id"])
            project_counts["found"] += 1
            status_ok = matched["status"] == label["expected_status"]
            if status_ok:
                project_counts["status_correct"] += 1
            else:
                kind = "错误判为有依据" if matched["status"] == "supported" and label["expected_status"] != "supported" else "状态判断错误"
                errors.append({"file": task["file"], "needle": label["needle"], "type": kind,
                               "detail": f"预期 {label['expected_status']}，实际 {matched['status']}"})
            if matched["status"] == "supported":
                project_counts["predicted_supported"] += 1
                if label["expected_status"] != "supported":
                    project_counts["false_supported"] += 1
            expected_ids = label["evidence_ids"]
            top = matched["evidence"][0] if matched["evidence"] else None
            evidence_ok = top["id"] in expected_ids if expected_ids else top is None
            project_counts["evidence_evaluated"] += 1
            if evidence_ok:
                project_counts["evidence_correct"] += 1
            else:
                errors.append({"file": task["file"], "needle": label["needle"], "type": "首要证据不符",
                               "detail": f"期望 {', '.join(expected_ids) or '无材料'}，实际 {top['id'] if top else '无材料'}"})
            section = sections.get(matched["id"], {})
            response = section.get("response", "")
            guardrail = {"supported": bool(top and section.get("evidence_name") == top["name"]),
                         "review": "待人工确认" in response,
                         "gap": "材料缺口" in response and "暂不作符合性承诺" in response,
                         "action": "待项目负责人跟进" in response}.get(matched["status"], False)
            if guardrail:
                project_counts["draft_guardrail_correct"] += 1
            else:
                errors.append({"file": task["file"], "needle": label["needle"], "type": "草稿保护失效",
                               "detail": "草稿未标明依据或待处理状态"})
            if status_ok and evidence_ok and guardrail:
                project_counts["end_to_end_correct"] += 1
        extra = [r for r in predictions if r["id"] not in used]
        project_counts["extracted"] = len(predictions)
        for row in extra:
            errors.append({"file": task["file"], "needle": row["text"], "type": "额外提取", "detail": "标注集中无对应要求"})
        for key, value in project_counts.items():
            counts[key] += value
            group[key] += value
        per_project.append({"file": task["file"], "split": split, "gold": project_counts["gold"],
                            "found": project_counts["found"], "status_correct": project_counts["status_correct"],
                            "evidence_correct": project_counts["evidence_correct"],
                            "end_to_end_correct": project_counts["end_to_end_correct"],
                            "false_supported": project_counts["false_supported"]})
    priority = {"错误判为有依据": 0, "要求遗漏": 1, "首要证据不符": 2,
                "状态判断错误": 3, "草稿保护失效": 4, "额外提取": 5}
    errors.sort(key=lambda row: priority.get(row["type"], 9))
    ratio = lambda numerator, denominator: round(numerator / denominator, 3) if denominator else 0
    return {"run_at": datetime.now().isoformat(timespec="seconds"), "mode": "离线规则基线",
            "project_count": len(tasks), "gold_count": counts["gold"], "extracted_count": counts["extracted"],
            "matched_count": counts["found"], "correct_status_count": counts["status_correct"],
            "evidence_correct_count": counts["evidence_correct"], "evidence_evaluated_count": counts["evidence_evaluated"],
            "end_to_end_correct_count": counts["end_to_end_correct"],
            "predicted_supported_count": counts["predicted_supported"], "false_supported_count": counts["false_supported"],
            "draft_guardrail_correct_count": counts["draft_guardrail_correct"],
            "requirement_recall": ratio(counts["found"], counts["gold"]),
            "requirement_precision": ratio(counts["found"], counts["extracted"]),
            "status_accuracy": ratio(counts["status_correct"], counts["found"]),
            "evidence_top1_accuracy": ratio(counts["evidence_correct"], counts["evidence_evaluated"]),
            "end_to_end_accuracy": ratio(counts["end_to_end_correct"], counts["gold"]),
            "false_supported_rate": ratio(counts["false_supported"], counts["predicted_supported"]),
            "draft_guardrail_rate": ratio(counts["draft_guardrail_correct"], counts["found"]),
            "elapsed_ms": round((time.perf_counter() - started) * 1000),
            "splits": [{"name": name, "gold": row["gold"], "found": row["found"],
                        "end_to_end_correct": row["end_to_end_correct"],
                        "false_supported": row["false_supported"]} for name, row in per_split.items()],
            "error_count": len(errors), "errors": errors, "projects": per_project,
            "note": "内置标注集上的离线规则基线；草稿保护只检查模板标记，不等同人工事实核验。"}
